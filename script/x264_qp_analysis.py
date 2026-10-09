#!/usr/bin/env python3
"""Run x264 for QP 0..51 and build an Excel report of file size and encoding speed.

Usage:
    python x264_qp_analysis.py foreman-cif.yuv --resolution 352x288
    python x264_qp_analysis.py clip.yuv --resolution 1280x720 --repeat 3 --x264 /path/to/x264

The report (CSV + XLSX with charts) is generated automatically once all runs finish.
Use --report-only to rebuild the report from an existing results.csv.
"""
import argparse
import csv
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = ["qp", "file_size_bytes", "fps", "time_sec", "bitrate_kbps"]
FPS_RE = re.compile(r"encoded (\d+) frames, ([\d.]+) fps, ([\d.]+) kb/s")


def run_x264(x264, yuv, resolution, qp, out_file, extra_args):
    """Encode once; return (frames, fps, bitrate_kbps, wall_time_sec, size_bytes)."""
    cmd = [x264, "--qp", str(qp), "--input-res", resolution, "--no-progress",
           *extra_args, "-o", str(out_file), str(yuv)]
    start = time.perf_counter()
    proc = subprocess.run(cmd, capture_output=True, text=True)
    elapsed = time.perf_counter() - start
    if proc.returncode != 0:
        raise RuntimeError(f"x264 failed for QP {qp}:\n{proc.stderr}")
    # x264 prints its summary on stderr
    m = FPS_RE.search(proc.stderr)
    if not m:
        raise RuntimeError(f"Could not parse x264 summary for QP {qp}:\n{proc.stderr}")
    frames, fps, kbps = int(m[1]), float(m[2]), float(m[3])
    return frames, fps, kbps, elapsed, out_file.stat().st_size


def run_all(args, outdir):
    rows = []
    enc_dir = outdir / "encoded"
    enc_dir.mkdir(parents=True, exist_ok=True)
    for qp in range(args.qp_min, args.qp_max + 1):
        runs = []
        for _ in range(args.repeat):
            out_file = enc_dir / f"qp{qp:02d}.264"
            runs.append(run_x264(args.x264, args.input, args.resolution, qp, out_file, args.x264_args))
        # Size is deterministic; speed is averaged over repeats
        n = len(runs)
        row = {
            "qp": qp,
            "file_size_bytes": runs[-1][4],
            "fps": round(sum(r[1] for r in runs) / n, 2),
            "time_sec": round(sum(r[3] for r in runs) / n, 3),
            "bitrate_kbps": round(sum(r[2] for r in runs) / n, 2),
        }
        rows.append(row)
        print(f"QP {qp:2d}: {row['file_size_bytes']:>10,} B  {row['fps']:>8.2f} fps  {row['time_sec']:.2f} s")
        if not args.keep_output:
            out_file.unlink(missing_ok=True)
    if not args.keep_output:
        shutil.rmtree(enc_dir, ignore_errors=True)
    return rows


def write_csv(rows, path):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def read_csv(path):
    with open(path, newline="") as f:
        return [{"qp": int(r["qp"]), "file_size_bytes": int(r["file_size_bytes"]),
                 "fps": float(r["fps"]), "time_sec": float(r["time_sec"]),
                 "bitrate_kbps": float(r["bitrate_kbps"])} for r in csv.DictReader(f)]


def add_chart(ws, title, y_title, col, anchor, n):
    ch = LineChart()
    ch.title, ch.x_axis.title, ch.y_axis.title = title, "QP", y_title
    ch.height, ch.width = 8.5, 16
    ch.legend = None
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    ws.add_chart(ch, anchor)


def write_xlsx(rows, path, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = ["QP", "File size (bytes)", "File size (KiB)", "Size vs QP0 (%)",
               "FPS", "Time (s)", "Bitrate (kb/s)", "Compression ratio"]
    ws.append(headers)
    n = len(rows)
    for i, r in enumerate(rows, start=2):
        ws.append([
            r["qp"], r["file_size_bytes"],
            f"=B{i}/1024",
            f"=B{i}/$B$2*100",
            r["fps"], r["time_sec"], r["bitrate_kbps"],
            f"={meta['raw_bytes']}/B{i}" if meta["raw_bytes"] else None,
        ])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0.0"
        row[3].number_format = "0.00"
        row[7].number_format = "0.0"
    ws.freeze_panes = "A2"

    add_chart(ws, "File size vs QP", "File size (KiB)", 3, "J2", n)
    add_chart(ws, "Encoding speed vs QP", "FPS", 5, "J20", n)
    add_chart(ws, "Encoding time vs QP", "Time (s)", 6, "J38", n)

    s = wb.create_sheet("Summary")
    last = n + 1
    summary = [
        ("Input", meta["input"]),
        ("Resolution", meta["resolution"]),
        ("Runs per QP (speed averaged)", meta["repeat"]),
        ("", ""),
        ("Largest file (bytes)", f"=MAX(Results!B2:B{last})"),
        ("Smallest file (bytes)", f"=MIN(Results!B2:B{last})"),
        ("Size reduction QP min -> QP max (x)", f"=Results!B2/Results!B{last}"),
        ("Slowest FPS", f"=MIN(Results!E2:E{last})"),
        ("Fastest FPS", f"=MAX(Results!E2:E{last})"),
        ("Speed-up QP min -> QP max (x)", f"=Results!E{last}/Results!E2"),
        ("Avg size change per +1 QP (%)", f"=(Results!B{last}/Results!B2)^(1/{n - 1})*100-100"),
    ]
    for row in summary:
        s.append(list(row))
    s.column_dimensions["A"].width = 40
    s.column_dimensions["B"].width = 24
    for c in s["A"]:
        c.font = Font(bold=True)
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("input", type=Path, help="uncompressed .yuv file")
    p.add_argument("--resolution", required=True, help="WxH, e.g. 352x288")
    p.add_argument("--x264", default=shutil.which("x264") or "x264", help="path to x264 executable")
    p.add_argument("--qp-min", type=int, default=0)
    p.add_argument("--qp-max", type=int, default=51)
    p.add_argument("--repeat", type=int, default=1, help="encodes per QP; FPS/time are averaged")
    p.add_argument("--outdir", type=Path, default=Path("results"))
    p.add_argument("--keep-output", action="store_true", help="keep the .264 files")
    p.add_argument("--report-only", action="store_true", help="rebuild report from results.csv")
    p.add_argument("x264_args", nargs="*", help="extra x264 args after '--', e.g. -- --threads 1")
    args = p.parse_args()

    if not re.fullmatch(r"\d+x\d+", args.resolution):
        sys.exit("--resolution must look like 352x288")
    if not args.input.is_file():
        sys.exit(f"Input not found: {args.input}")
    if not args.report_only and not shutil.which(args.x264):
        sys.exit(f"x264 not found: {args.x264}")
    args.outdir.mkdir(parents=True, exist_ok=True)
    csv_path = args.outdir / "results.csv"

    if args.report_only:
        rows = read_csv(csv_path)
    else:
        rows = run_all(args, args.outdir)
        write_csv(rows, csv_path)

    meta = {"input": args.input.name, "resolution": args.resolution, "repeat": args.repeat,
            "raw_bytes": args.input.stat().st_size}
    xlsx_path = args.outdir / "qp_report.xlsx"
    write_xlsx(rows, xlsx_path, meta)
    print(f"\nWrote {csv_path} and {xlsx_path}")


if __name__ == "__main__":
    main()
